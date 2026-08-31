#!/usr/bin/env python3
"""Parse Handelsbanken xlsx exports -> unified transactions CSV.

Usage: python3 scripts/bank-parse.py finance/data/*.xlsx --out finance/data/transactions.csv

Month-to-month exports typically overlap; rows repeated across files are
dropped automatically (keyed on account, date, description, amount, balance).
Pass --keep-duplicates to write every raw row instead.
"""
import argparse, csv, re, sys
from pathlib import Path

def parse_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        # find header row: Reskontradatum
        hdr_i = next((i for i, r in enumerate(rows) if r and r[0] == "Reskontradatum"), None)
        if hdr_i is None:
            continue
        # account name: scan above header for the line that is the account title (not Saldo/Period/Kontoform)
        account = None
        for r in rows[:hdr_i]:
            for c in r:
                if c and isinstance(c, str) and re.search(r"\d{3} \d{3} \d{3}", c):
                    account = c.strip()
                    break
            if account: break
        for r in rows[hdr_i + 1:]:
            if not r or not r[0]: continue
            out.append({
                "account": account or path.stem,
                "reskontra": str(r[0]) if r[0] else "",
                "trans_date": str(r[1]) if r[1] else "",
                "text": str(r[2]).strip() if r[2] else "",
                "amount": float(r[3]) if r[3] is not None else 0.0,
                "balance": float(r[4]) if r[4] is not None else None,
            })
    return out

def parse_ica_csv(path: Path) -> list[dict]:
    """Parse ICA 'Kontohändelser' exports: ;-separated, Swedish decimal comma.

    Header: Datum;Text;Typ;Belopp;Saldo. Amounts like '-4,00 kr', balances
    like '3 075,52 kr'. Belopp is the transaction amount, not Importsaldo.
    """
    def num(s: str) -> float:
        if not s: return 0.0
        s = s.replace(" ", "").replace("kr", "").strip().replace(",", ".")
        try: return float(s)
        except ValueError: return 0.0
    out = []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        for r in csv.reader(fh, delimiter=";"):
            if not r or not r[0] or r[0].lower() == "datum": continue
            if len(r) < 4: continue
            out.append({
                "account": "ICA Matkonto",
                "reskontra": "",
                "trans_date": r[0].strip(),
                "text": r[1].strip() + (" (" + r[2].strip() + ")" if len(r) > 2 and r[2] else ""),
                "amount": num(r[3]),
                "balance": num(r[4]) if len(r) > 4 else None,
            })
    return out

def dedupe(rows: list[dict]) -> list[dict]:
    """Drop rows repeated across overlapping exports.

    Keyed on (account, trans_date, text, amount, balance): overlapping
    month-to-month exports repeat rows identically on every field, while two
    genuinely distinct transactions that share date/description/amount (e.g.
    two equal card purchases the same day) differ in the running balance and
    are both kept.
    """
    seen, uniq = set(), []
    for r in rows:
        k = (r["account"], r["trans_date"], r["text"], r["amount"], r["balance"])
        if k in seen:
            continue
        seen.add(k)
        uniq.append(r)
    return uniq

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--out", default="finance/data/transactions.csv")
    ap.add_argument("--keep-duplicates", action="store_true",
                    help="write every raw row, skipping the overlap dedupe pass")
    args = ap.parse_args()
    rows = []
    for f in args.files:
        p = Path(f)
        if p.suffix.lower() == ".csv":
            rows += parse_ica_csv(p)
        else:
            rows += parse_xlsx(p)
    if args.keep_duplicates:
        kept = rows
    else:
        kept = dedupe(rows)
        dropped = len(rows) - len(kept)
        if dropped:
            print(f"note: dropped {dropped} duplicate row(s) from overlapping exports "
                  f"(use --keep-duplicates to keep them)", file=sys.stderr)
    kept.sort(key=lambda r: (r["trans_date"], r["account"]))
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["account", "reskontra", "trans_date", "text", "amount", "balance"])
        w.writeheader(); w.writerows(kept)
    print(f"{len(kept)} transactions ({len(rows)} raw) -> {out}")

if __name__ == "__main__":
    main()
